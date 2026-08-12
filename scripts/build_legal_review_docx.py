from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "legal-review"
OUT = SOURCE / "Mesa-OS-Pacote-Revisao-Juridica-Contexto-Longitudinal-v0.1.docx"
FILES = ["README.md", "01-TERMS-OF-USE-DRAFT.md", "02-PRIVACY-NOTICE-DRAFT.md", "03-TUTORIA-LONGITUDINAL-CONTEXT-NOTICE-DRAFT.md", "04-LEGAL-REVIEW-CHECKLIST.md"]

def font(run, size=11, bold=False, color="101D37"):
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)

def paragraph(doc, text="", size=11, bold=False, color="222222", after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(text)
    font(r, size, bold, color)
    return p

doc = Document()
section = doc.sections[0]
section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Inches(1)
for style_name, size, color, before, after in [("Normal", 11, "222222", 0, 6), ("Heading 1", 16, "101D37", 16, 8), ("Heading 2", 13, "2E74B5", 12, 6), ("Heading 3", 12, "101D37", 8, 4)]:
    style = doc.styles[style_name]
    style.font.name = "Arial"; style._element.rPr.rFonts.set(qn("w:ascii"), "Arial"); style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    style.font.size = Pt(size); style.font.color.rgb = RGBColor.from_string(color)
    style.paragraph_format.space_before = Pt(before); style.paragraph_format.space_after = Pt(after); style.paragraph_format.line_spacing = 1.15

p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before = Pt(150); p.paragraph_format.space_after = Pt(8)
font(p.add_run("MESA DOS DONOS"), 13, True, "101D37")
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(12)
font(p.add_run("Pacote para revisão jurídica"), 26, True, "101D37")
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(20)
font(p.add_run("Contexto Longitudinal do TutorIA"), 16, False, "444444")
paragraph(doc, "Versão de trabalho 0.1 · 12 de agosto de 2026", 10, False, "666666", 6).alignment = WD_ALIGN_PARAGRAPH.CENTER
paragraph(doc, "Material de apoio para análise jurídica. Não publicar nem ativar coleta automática antes da revisão e aprovação formal.", 10, False, "9B1C1C", 0).alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.add_page_break()

for filename in FILES:
    lines = (SOURCE / filename).read_text(encoding="utf-8").splitlines()
    for line in lines:
        stripped = line.strip()
        if not stripped:
            continue
        if stripped.startswith("# "):
            doc.add_heading(stripped[2:], level=1)
        elif stripped.startswith("## "):
            doc.add_heading(stripped[3:], level=2)
        elif stripped.startswith("### "):
            doc.add_heading(stripped[4:], level=3)
        elif stripped.startswith("- [ ] "):
            p = doc.add_paragraph(style="List Bullet"); p.paragraph_format.space_after = Pt(4); font(p.add_run("☐ " + stripped[6:]), 11, False, "222222")
        elif stripped.startswith("- "):
            p = doc.add_paragraph(style="List Bullet"); p.paragraph_format.space_after = Pt(4); font(p.add_run(stripped[2:]), 11, False, "222222")
        elif len(stripped) > 3 and stripped[0].isdigit() and ". " in stripped[:4]:
            p = doc.add_paragraph(style="List Number"); p.paragraph_format.space_after = Pt(4); font(p.add_run(stripped.split(". ", 1)[1]), 11, False, "222222")
        elif stripped.startswith("**") and stripped.endswith("**"):
            paragraph(doc, stripped.strip("*"), 11, True, "101D37", 4)
        else:
            clean = stripped.replace("**", "").replace("`", "")
            paragraph(doc, clean)
    doc.add_page_break()

for section in doc.sections:
    footer = section.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    font(footer.add_run("Mesa OS · Pacote para revisão jurídica · v0.1"), 8, False, "777777")

doc.core_properties.title = "Pacote para revisão jurídica — Contexto Longitudinal TutorIA"
doc.core_properties.author = "Mesa dos Donos"
doc.core_properties.subject = "Rascunho para revisão jurídica"
doc.save(OUT)
print(OUT)
